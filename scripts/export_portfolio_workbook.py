"""Export a multi-sheet portfolio tracking workbook from local canonical data."""
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

ROOT = Path(__file__).resolve().parent.parent
DATA_DIR = ROOT / "data"
OUTPUT = ROOT / "Portfolio_Tracker.xlsx"


def read_csv(name, columns=None):
    path = DATA_DIR / name
    if not path.exists() or not path.stat().st_size:
        return pd.DataFrame(columns=columns or [])
    return pd.read_csv(path, dtype=str).fillna("")


def number(value):
    try:
        return float(value or 0)
    except (TypeError, ValueError):
        return 0.0


def build_summary(holdings, transactions, income):
    numeric = holdings.copy()
    for column in ("Inv. Amount", "Latest Value", "Overall Gain", "Day's Gain"):
        if column in numeric.columns:
            numeric.loc[:, column] = pd.to_numeric(numeric[column], errors="coerce").fillna(0)
    invested = numeric["Inv. Amount"].sum()
    current = numeric["Latest Value"].sum()
    gain = current - invested
    income_total = sum(number(row.get("Total Amount")) for row in income.to_dict(orient="records"))
    buys = transactions[transactions.get("Type", pd.Series(dtype=str)).eq("BUY")] if not transactions.empty else transactions
    sells = transactions[transactions.get("Type", pd.Series(dtype=str)).eq("SELL")] if not transactions.empty else transactions
    return pd.DataFrame([
        ["Report generated", datetime.now().astimezone().isoformat(timespec="seconds")],
        ["Active lots", len(holdings)],
        ["Total invested", invested],
        ["Current value", current],
        ["Unrealized gain", gain],
        ["Unrealized return %", gain / invested if invested else 0],
        ["Tracked income", income_total],
        ["Buy transactions", len(buys)],
        ["Sell transactions", len(sells)],
    ], columns=["Metric", "Value"])


def build_income():
    dividends = read_csv("dividends.csv")
    if dividends.empty:
        return dividends
    dividends["Total Amount"] = dividends.apply(
        lambda row: number(row.get("Total Amount")) or number(row.get("Amount Per Share")) * number(row.get("Quantity Eligible")), axis=1
    )
    return dividends


def add_sheet_table(ws, frame, table_name):
    if frame.empty:
        ws.append(["No records"])
        return
    for row in [list(frame.columns), *frame.astype(object).where(pd.notna(frame), "").values.tolist()]:
        ws.append(row)
    ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
    table = Table(displayName=table_name, ref=ref)
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False, showRowStripes=True, showColumnStripes=False)
    ws.add_table(table)


def format_sheet(ws):
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    thin = Side(style="thin", color="D9E2F3")
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(bottom=thin)
    for column_cells in ws.columns:
        values = [len(str(cell.value or "")) for cell in column_cells[:100]]
        width = min(max(max(values, default=10) + 2, 10), 34)
        ws.column_dimensions[get_column_letter(column_cells[0].column)].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    ws.row_dimensions[1].height = 24


def format_summary(ws):
    for row in range(2, ws.max_row + 1):
        if ws.cell(row, 1).value in {"Total invested", "Current value", "Unrealized gain", "Tracked income"}:
            ws.cell(row, 2).number_format = '₹#,##,##0.00'
        if ws.cell(row, 1).value == "Unrealized return %":
            ws.cell(row, 2).number_format = '0.00%'


def format_holdings(ws):
    headers = {cell.value: cell.column for cell in ws[1]}
    for name in ("Inv. Amount", "Latest Value", "Overall Gain", "Day's Gain", "Inv. Price", "Latest Price/NAV"):
        column = headers.get(name)
        if column:
            for cell in ws.iter_cols(min_col=column, max_col=column, min_row=2):
                for value_cell in cell:
                    value_cell.number_format = '₹#,##0.00'
    gain_column = headers.get("Overall Gain")
    if gain_column:
        letter = get_column_letter(gain_column)
        ws.conditional_formatting.add(f"{letter}2:{letter}{ws.max_row}", CellIsRule(operator="greaterThan", formula=["0"], font=Font(color="008000")))
        ws.conditional_formatting.add(f"{letter}2:{letter}{ws.max_row}", CellIsRule(operator="lessThan", formula=["0"], font=Font(color="C00000")))


def export_workbook(output=OUTPUT):
    holdings = read_csv("holdings.csv")
    transactions = read_csv("transactions.csv")
    corrections = read_csv("holding_corrections.csv")
    deletions = read_csv("holding_deletions.csv")
    income = build_income()
    distributions = read_csv("distribution_history.csv")
    summary = build_summary(holdings, transactions, income)

    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        summary.to_excel(writer, index=False, sheet_name="Summary")
        holdings.to_excel(writer, index=False, sheet_name="Holdings")
        transactions.to_excel(writer, index=False, sheet_name="Transactions")
        income.to_excel(writer, index=False, sheet_name="Income")
        distributions.to_excel(writer, index=False, sheet_name="Distribution History")
        corrections.to_excel(writer, index=False, sheet_name="Corrections")
        deletions.to_excel(writer, index=False, sheet_name="Deletions")

    workbook = load_workbook(output)
    for sheet in workbook.worksheets:
        format_sheet(sheet)
    format_summary(workbook["Summary"])
    format_holdings(workbook["Holdings"])
    workbook["Summary"].freeze_panes = "A2"
    workbook.save(output)
    print(f"Wrote portfolio tracker workbook: {output}")
    return output


if __name__ == "__main__":
    export_workbook()
